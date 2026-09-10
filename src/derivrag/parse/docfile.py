"""Word 文档解析（.doc / .docx）。

为什么这个模块不能省：SSE 有两批文档的正文只存在于 Word 附件里。

  - 试点业务规则 16 篇中有 6 篇，正文在 `.doc`（Word 97 OLE2，magic d0cf11e0）
  - 2025 规则库的文章页 HTML 壳只有约 1550 个可见字符，
    真正的规则条文完全在 `.docx` 附件中（页面用 mammoth.js 在前端渲染）

PyMuPDF 读不了 Word。python-docx 只能读 .docx，读不了 OLE2 的 .doc，
所以 .doc 需要先用 LibreOffice 无头转换成 .docx。
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .clean import clean_text, looks_chinese

logger = logging.getLogger(__name__)

# LibreOffice 的常见安装位置
_SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/local/bin/soffice",
    "/opt/homebrew/bin/soffice",
    "soffice",
    "libreoffice",
)


def find_soffice(configured: str | None = None) -> str | None:
    """定位 LibreOffice 可执行文件。找不到返回 None。"""
    candidates = ([configured] if configured else []) + list(_SOFFICE_CANDIDATES)
    for c in candidates:
        if not c:
            continue
        if Path(c).is_file():
            return c
        found = shutil.which(c)
        if found:
            return found
    return None


def doc_to_docx(path: Path, *, soffice: str | None = None, timeout: int = 180) -> Path | None:
    """用 LibreOffice 把 .doc 转成 .docx，返回新文件路径。

    转换结果写到临时目录，调用方负责用完即弃。
    """
    exe = find_soffice(soffice)
    if not exe:
        logger.warning(
            "未找到 LibreOffice，无法解析 %s。"
            "安装方式: brew install --cask libreoffice",
            path.name,
        )
        return None

    outdir = Path(tempfile.mkdtemp(prefix="derivrag_doc_"))
    try:
        subprocess.run(
            [exe, "--headless", "--convert-to", "docx", "--outdir", str(outdir), str(path)],
            check=True,
            capture_output=True,
            timeout=timeout,
        )
    except subprocess.TimeoutExpired:
        logger.warning("LibreOffice 转换 %s 超时", path.name)
        return None
    except subprocess.CalledProcessError as e:
        logger.warning("LibreOffice 转换 %s 失败: %s", path.name, e.stderr[:300])
        return None

    converted = outdir / f"{path.stem}.docx"
    return converted if converted.exists() else None


# =====================================================================
# Word 自动编号还原
# =====================================================================
#
# SSE 的规则 .doc 里，"第一条""第二条"这些条号**不是正文文本**，
# 而是 Word 的自动列表编号 —— 它只存在于 numbering.xml 的格式定义中，
# python-docx 的 paragraph.text 完全读不到。
#
# 实测 SSE《股票期权试点交易规则》：正文里能匹配到的"第X条"只有 10 处
# （全是"依照本规则第X条"这类交叉引用），而真实条数是 170 条。
# 不还原编号，这份文档就只能退化成按段落硬切，彻底失去"精确到条"的引用能力
# —— 而这恰恰是本项目最核心的价值。
#
# 还原方式不靠猜：直接读 word/numbering.xml，它明确写着
#     <w:numFmt w:val="chineseCountingThousand"/>
#     <w:lvlText w:val="第%1条"/>
# 于是按文档顺序对每个 numPr 段落计数，再套用这个模板即可完全复现
# Word 的显示结果。

_CN_NUM = "零一二三四五六七八九"
_CN_UNIT = ["", "十", "百", "千"]


def _int_to_chinese(n: int) -> str:
    """整数转中文数字（chineseCountingThousand 格式）。

    1->一  10->十  11->十一  20->二十  103->一百零三  110->一百一十
    """
    if n <= 0:
        return "零"
    if n >= 10000:
        return str(n)

    digits = [int(c) for c in str(n)]
    out: list[str] = []
    length = len(digits)
    zero_pending = False
    for i, d in enumerate(digits):
        unit = _CN_UNIT[length - i - 1]
        if d == 0:
            zero_pending = True
            continue
        if zero_pending and out:
            out.append(_CN_NUM[0])
        zero_pending = False
        # 中文习惯: 十一 而非 一十一（仅当十位是最高位时省略"一"）
        if d == 1 and unit == "十" and i == 0:
            out.append(unit)
        else:
            out.append(_CN_NUM[d] + unit)
    return "".join(out)


def _format_number(value: int, num_fmt: str) -> str:
    """按 Word 的 numFmt 渲染序号。"""
    if num_fmt in ("chineseCountingThousand", "chineseCounting", "taiwaneseCountingThousand",
                   "taiwaneseCounting", "japaneseCounting"):
        return _int_to_chinese(value)
    if num_fmt == "decimalEnclosedCircle":
        return "①②③④⑤⑥⑦⑧⑨⑩"[value - 1] if 1 <= value <= 10 else str(value)
    if num_fmt == "lowerLetter":
        return chr(ord("a") + (value - 1) % 26)
    if num_fmt == "upperLetter":
        return chr(ord("A") + (value - 1) % 26)
    if num_fmt == "lowerRoman":
        return _to_roman(value).lower()
    if num_fmt == "upperRoman":
        return _to_roman(value)
    if num_fmt == "none":
        return ""
    return str(value)


def _to_roman(n: int) -> str:
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
            (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def _read_numbering_defs(docx_path: Path) -> dict[str, dict[str, Any]]:
    """读 word/numbering.xml。

    返回 {numId: {"abstract": abstractNumId, "levels": {ilvl: {numFmt, lvlText, start}}}}

    必须把 abstractNumId 一并带出来：多个 numId 可以指向同一个
    abstractNum（本例中 numId 3 和 4 都指向 abstractNum 3），此时 Word
    是**共用同一个计数器**继续往下编号的。如果按 numId 分别计数，
    numId=4 的那一段会从"第一条"重新开始，正文里就会冒出两个"第一条"。
    """
    import re
    import zipfile

    defs: dict[str, dict[str, Any]] = {}
    try:
        with zipfile.ZipFile(docx_path) as z:
            if "word/numbering.xml" not in z.namelist():
                return defs
            xml = z.read("word/numbering.xml").decode("utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        return defs

    # numId -> abstractNumId
    num_map = dict(
        re.findall(
            r'<w:num\s+w:numId="(\d+)"[^>]*>\s*<w:abstractNumId\s+w:val="(\d+)"', xml
        )
    )

    abstract: dict[str, dict[int, dict[str, str]]] = {}
    for m in re.finditer(
        r'<w:abstractNum\s+w:abstractNumId="(\d+)".*?</w:abstractNum>', xml, re.S
    ):
        aid, block = m.group(1), m.group(0)
        levels: dict[int, dict[str, str]] = {}
        for lm in re.finditer(r'<w:lvl\s+w:ilvl="(\d+)".*?</w:lvl>', block, re.S):
            ilvl, lblock = int(lm.group(1)), lm.group(0)
            fmt = re.search(r'<w:numFmt\s+w:val="([^"]+)"', lblock)
            text = re.search(r'<w:lvlText\s+w:val="([^"]*)"', lblock)
            start = re.search(r'<w:start\s+w:val="(\d+)"', lblock)
            levels[ilvl] = {
                "numFmt": fmt.group(1) if fmt else "decimal",
                "lvlText": text.group(1) if text else "",
                "start": start.group(1) if start else "1",
            }
        abstract[aid] = levels

    for num_id, abs_id in num_map.items():
        if abs_id in abstract:
            defs[num_id] = {"abstract": abs_id, "levels": abstract[abs_id]}
    return defs


def _paragraph_numbering(paragraph) -> tuple[str, int] | None:
    """取段落的 (numId, ilvl)，无自动编号时返回 None。"""
    from docx.oxml.ns import qn

    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        return None
    num_id_el = numpr.find(qn("w:numId"))
    ilvl_el = numpr.find(qn("w:ilvl"))
    if num_id_el is None:
        return None
    num_id = num_id_el.get(qn("w:val"))
    ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
    return (num_id, ilvl) if num_id else None


class _NumberingRenderer:
    """按文档顺序复现 Word 的自动编号。

    计数器以 **abstractNumId** 为键，与 Word 的行为一致 —— 共用同一个
    abstractNum 的多个 numId 共享一条编号序列，不会各自从头开始。
    """

    def __init__(self, defs: dict[str, dict[str, Any]]) -> None:
        self.defs = defs
        self.counters: dict[tuple[str, int], int] = {}

    def label_for(self, num_id: str, ilvl: int) -> str:
        entry = self.defs.get(num_id)
        if not entry:
            return ""
        levels = entry["levels"]
        if ilvl not in levels:
            return ""
        spec = levels[ilvl]
        abs_id = entry["abstract"]

        key = (abs_id, ilvl)
        if key not in self.counters:
            self.counters[key] = int(spec.get("start", "1"))
        else:
            self.counters[key] += 1
        # 进入新的上层编号时，重置所有更深层级
        for k in list(self.counters):
            if k[0] == abs_id and k[1] > ilvl:
                del self.counters[k]

        lvl_text = spec.get("lvlText", "")
        if not lvl_text:
            return ""
        # lvlText 里的 %1 %2 … 指代第 1、2 … 层的当前计数
        out = lvl_text
        for lv in range(ilvl + 1):
            cnt = self.counters.get((abs_id, lv))
            if cnt is None:
                continue
            fmt = levels.get(lv, {}).get("numFmt", "decimal")
            out = out.replace(f"%{lv + 1}", _format_number(cnt, fmt))
        return out


def _iter_block_items(doc):
    """按文档顺序依次产出段落与表格。

    python-docx 的 doc.paragraphs / doc.tables 是两个独立列表，
    直接用会丢失段落与表格的相对位置。这里走底层 XML 保序遍历，
    否则条款正文和它下面的表格会被拆散。
    """
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = doc.element.body if isinstance(doc, _Doc) else doc
    for child in parent.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def parse_docx(path: str | Path) -> dict[str, Any]:
    """解析 .docx，保持段落与表格的原始顺序。"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    path = Path(path)
    document = docx.Document(str(path))

    # 还原 Word 自动编号（第X条 等），否则条款切分无从下手
    renderer = _NumberingRenderer(_read_numbering_defs(path))

    parts: list[str] = []
    tables: list[dict[str, Any]] = []

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            txt = block.text.strip()
            numbering = _paragraph_numbering(block)
            if numbering is not None:
                label = renderer.label_for(*numbering)
                if label:
                    # Word 的编号与正文之间没有分隔符，直接前缀即可，
                    # 与 PDF 版规则里"第一条为规范……"的排版完全一致
                    txt = f"{label}{txt}" if txt else label
            if txt:
                parts.append(txt)
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells = [clean_text(c.text, is_cjk=True) for c in row.cells]
                # 合并单元格会让同一个值重复出现，去掉相邻重复
                dedup = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                if any(dedup):
                    rows.append(dedup)
            if rows:
                idx = len(tables)
                tables.append(
                    {
                        "index": idx,
                        "n_rows": len(rows),
                        "n_cols": max(len(r) for r in rows),
                        "rows": rows,
                    }
                )
                # 表格按原位置渲染进正文
                lines = []
                for row in rows:
                    cells = [c for c in row if c]
                    if len(cells) == 2:
                        lines.append(f"{cells[0]}：{cells[1]}")
                    elif cells:
                        lines.append(" | ".join(cells))
                parts.append("\n".join(lines))

    raw = "\n\n".join(parts)
    is_cjk = looks_chinese(raw)
    text = clean_text(raw, is_cjk=is_cjk)

    warnings: list[str] = []
    if not text.strip():
        warnings.append("Word 文档解析结果为空")

    return {
        "text": text,
        "tables": tables,
        "page_count": None,
        "ocr_pages": [],
        "warnings": warnings,
        "lang": "zh" if is_cjk else "en",
    }


def parse_doc(path: str | Path, *, soffice: str | None = None) -> dict[str, Any]:
    """解析 .doc —— 先转 .docx 再走 parse_docx。"""
    path = Path(path)
    converted = doc_to_docx(path, soffice=soffice)
    if converted is None:
        return {
            "text": "",
            "tables": [],
            "page_count": None,
            "ocr_pages": [],
            "warnings": [
                f"{path.name} 是 Word 97 (.doc) 格式，需要 LibreOffice 转换但转换未成功。"
                "安装: brew install --cask libreoffice"
            ],
            "lang": "zh",
        }
    try:
        return parse_docx(converted)
    finally:
        shutil.rmtree(converted.parent, ignore_errors=True)
